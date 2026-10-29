"""Who competes for and wins the kind of tenders the company targets.

Reads data/historico.json (from download_history.py) and writes data/competencia.{json,md,xlsx,html}:
  - summary by size bracket: tenders, share awarded / deserted, bidders per tender, awarded amounts
  - competitors: bids, wins, win rate, amount won, main buyers, per supplier (by RUT)
  - incumbents: suppliers that win the same buyer's tenders again and again
  - for each tender in the current ranking: past tenders of the same buyer, or with the same
    Código ONU and size bracket, with their winners, bidder counts and awarded amounts

Amounts are line totals awarded (MontoLineaAdjudica), converted to CLP at approximate rates.

Usage:
    python competitors.py [--historico data/historico.json] [--profile data/company_profile.json]
                          [--ranking data/ranking_top100.json] [--data data/licitaciones.json]
                          [--out data/competencia] [--open]
"""
import argparse
import collections
import json
import pathlib
import statistics
import sys

SIZE_ORDER = ["L1", "LE", "LP", "LQ", "LR"]
MIN_INCUMBENT_WINS = 2


def buyer_key(code):
    # The first part of the tender code identifies the buying unit (1867-50-LE26 -> 1867).
    return code.split("-", 1)[0]


def median(xs):
    xs = [x for x in xs if x]
    return statistics.median(xs) if xs else 0


def mclp(x):
    return round(x / 1e6, 1)


def top_names(counter, n=3):
    return ", ".join(f"{k} ({v})" for k, v in counter.most_common(n))


def winners(t):
    return [o for o in t["ofertas"] if o["gano"]]


def awarded_amount(t):
    return sum(o["adjudicado"] for o in winners(t))


def summarize_sizes(tenders):
    rows = []
    groups = collections.defaultdict(list)
    for t in tenders:
        groups[t["tamano"] if t["tamano"] in SIZE_ORDER else "otro"].append(t)
    for size in SIZE_ORDER + ["otro"]:
        ts = groups.get(size)
        if not ts:
            continue
        won = [t for t in ts if winners(t)]
        bidders = [t["n_oferentes"] for t in won if t["n_oferentes"]]
        rows.append({
            "tamano": size,
            "licitaciones": len(ts),
            "adjudicadas": len(won),
            "desiertas": sum(1 for t in ts if t["estado"].lower().startswith("desierta")),
            "oferentes_promedio": round(statistics.mean(bidders), 1) if bidders else None,
            "pct_un_oferente": round(100 * sum(1 for b in bidders if b == 1) / len(bidders)) if bidders else None,
            "monto_mediano_mm": mclp(median(awarded_amount(t) for t in won)),
        })
    return rows


def supplier_stats(tenders, sizes):
    s = {}
    for t in tenders:
        decided = bool(winners(t))
        in_segment = t["tamano"] in sizes
        for o in t["ofertas"]:
            p = s.setdefault(o["rut"], {"rut": o["rut"], "names": collections.Counter(), "ofertadas": 0,
                                         "ofertadas_resueltas": 0, "ganadas": 0, "monto_ganado": 0.0,
                                         "ofertadas_segmento": 0, "ganadas_segmento": 0,
                                         "compradores": collections.Counter(), "regiones": collections.Counter(),
                                         "montos": [], "ultima": "", "ejemplos": []})
            p["names"][o["nombre"]] += 1
            p["ofertadas"] += 1
            p["ofertadas_resueltas"] += decided
            p["ofertadas_segmento"] += in_segment
            if o["gano"]:
                p["ganadas"] += 1
                p["ganadas_segmento"] += in_segment
                p["monto_ganado"] += o["adjudicado"]
                p["montos"].append(o["adjudicado"])
                p["compradores"][t["organismo"]] += 1
                p["regiones"][t["region"].replace("Región ", "").replace("de ", "", 1).strip()] += 1
                p["ultima"] = max(p["ultima"], t["fecha_adjudicacion"] or t["fecha_publicacion"])
                if len(p["ejemplos"]) < 3:
                    p["ejemplos"].append(f"{t['codigo']} {t['nombre'][:60]}")
    out = []
    for p in s.values():
        out.append({
            "rut": p["rut"],
            "nombre": p["names"].most_common(1)[0][0],
            "ofertadas": p["ofertadas"],
            "ganadas": p["ganadas"],
            "tasa_exito": round(100 * p["ganadas"] / p["ofertadas_resueltas"]) if p["ofertadas_resueltas"] else None,
            "monto_ganado_mm": mclp(p["monto_ganado"]),
            "monto_mediano_mm": mclp(median(p["montos"])),
            "ofertadas_segmento": p["ofertadas_segmento"],
            "ganadas_segmento": p["ganadas_segmento"],
            "compradores": top_names(p["compradores"]),
            "regiones": top_names(p["regiones"], 2),
            "ultima_adjudicacion": p["ultima"],
            "ejemplos": " | ".join(p["ejemplos"]),
        })
    out.sort(key=lambda p: (-p["ganadas"], -p["monto_ganado_mm"]))
    return out


def incumbents(tenders):
    by_buyer = collections.defaultdict(list)
    for t in tenders:
        if winners(t):
            by_buyer[buyer_key(t["codigo"])].append(t)
    out = []
    for key, ts in by_buyer.items():
        wins = collections.Counter()
        names = {}
        for t in ts:
            for o in winners(t):
                wins[o["rut"]] += 1
                names[o["rut"]] = o["nombre"]
        rut, n = wins.most_common(1)[0]
        if n < MIN_INCUMBENT_WINS:
            continue
        out.append({
            "comprador": ts[0]["organismo"],
            "unidad": ts[0]["unidad"],
            "codigo_comprador": key,
            "licitaciones_adjudicadas": len(ts),
            "incumbente": names[rut],
            "rut": rut,
            "ganadas": n,
            "pct": round(100 * n / len(ts)),
            "otros_ganadores": ", ".join(f"{names[r]} ({c})" for r, c in wins.most_common(4)[1:]),
            "ejemplos": " | ".join(f"{t['codigo']} {t['nombre'][:50]}" for t in ts[:3]),
        })
    out.sort(key=lambda r: (-r["ganadas"], -r["pct"]))
    return out


def benchmark(rank_row, items_onu, by_buyer, tenders):
    key = buyer_key(rank_row["codigo"])
    past = [t for t in by_buyer.get(key, []) if t["codigo"] != rank_row["codigo"]]
    basis = "mismo comprador"
    if not [t for t in past if winners(t)]:
        past = []
        for digits in (8, 6):
            codes = {c[:digits] for c in items_onu if len(c) >= digits}
            if not codes:
                continue
            past = [t for t in tenders if t["tamano"] == rank_row["tamano"]
                    and any(c[:digits] in codes for c in t["codigos_onu"])]
            if [t for t in past if winners(t)]:
                basis = f"mismo código ONU ({digits} díg.) y tamaño"
                break
        else:
            basis = ""
    won = [t for t in past if winners(t)]
    wins = collections.Counter(o["nombre"] for t in won for o in winners(t))
    bidders = [t["n_oferentes"] for t in won if t["n_oferentes"]]
    # A buyer's past LR contract says little about the price of its LE tender.
    same_size = [t for t in won if t["tamano"] == rank_row["tamano"]]
    return {
        "rank": rank_row["rank"],
        "codigo": rank_row["codigo"],
        "nombre": rank_row["nombre"],
        "organismo": rank_row["organismo"],
        "tamano": rank_row["tamano"],
        "url": rank_row.get("url", ""),
        "base": basis if won else "sin historial comparable",
        "previas": len(past),
        "previas_adjudicadas": len(won),
        "ganadores_previos": top_names(wins, 4),
        "oferentes_promedio": round(statistics.mean(bidders), 1) if bidders else None,
        "monto_mediano_mm": mclp(median(awarded_amount(t) for t in same_size)) if same_size else None,
        "previas_codigos": [t["codigo"] for t in sorted(won, key=lambda t: t["fecha_publicacion"], reverse=True)[:12]],
        "ejemplos": " | ".join(f"{t['codigo']} {t['nombre'][:45]} → "
                               f"{', '.join(o['nombre'][:25] for o in winners(t)) or t['estado']}"
                               for t in won[:3]),
    }


def cell(s, n=None):
    s = "" if s is None else str(s)
    s = s.replace("|", "/").replace("\n", " ")
    return s[:n] + "…" if n and len(s) > n else s


def md_table(rows, cols, widths=None):
    widths = widths or {}
    lines = ["| " + " | ".join(h for h, _ in cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        lines.append("| " + " | ".join(cell(r.get(k), widths.get(k)) for _, k in cols) + " |")
    return "\n".join(lines)


def write_md(path, res):
    m = res["meta"]
    seg = ", ".join(m["tamanos_perfil"]) or "todos"
    parts = [
        f"# Competencia en Mercado Público — {m['empresa']}",
        f"Historial {m['meses']} · {m['licitaciones']} licitaciones relevantes para el perfil "
        f"({m['adjudicadas']} adjudicadas) · {m['proveedores']} proveedores distintos ofertaron. "
        "Montos en millones de CLP (conversión aproximada para UF/USD/EUR).",
        "## Por tamaño", md_table(res["por_tamano"], [
            ("Tamaño", "tamano"), ("Licitaciones", "licitaciones"), ("Adjudicadas", "adjudicadas"),
            ("Desiertas", "desiertas"), ("Oferentes prom.", "oferentes_promedio"),
            ("% con 1 oferente", "pct_un_oferente"), ("Monto mediano (MM)", "monto_mediano_mm")]),
        f"## Competidores en tu segmento ({seg})",
        md_table(res["competidores_segmento"][:25], [
            ("RUT", "rut"), ("Proveedor", "nombre"), ("Ganadas segm.", "ganadas_segmento"),
            ("Ofertadas segm.", "ofertadas_segmento"), ("Ganadas total", "ganadas"), ("Tasa %", "tasa_exito"),
            ("Ganado (MM)", "monto_ganado_mm"), ("Compradores", "compradores"), ("Regiones", "regiones")],
            {"nombre": 40, "compradores": 70}),
        "## Mayores ganadores (todos los tamaños)",
        md_table(res["proveedores"][:20], [
            ("RUT", "rut"), ("Proveedor", "nombre"), ("Ganadas", "ganadas"), ("Ofertadas", "ofertadas"),
            ("Tasa %", "tasa_exito"), ("Ganado (MM)", "monto_ganado_mm"), ("Compradores", "compradores")],
            {"nombre": 40, "compradores": 70}),
        "## Incumbentes por comprador",
        md_table(res["incumbentes"][:30], [
            ("Comprador", "comprador"), ("Adjudicadas", "licitaciones_adjudicadas"), ("Incumbente", "incumbente"),
            ("Ganadas", "ganadas"), ("%", "pct"), ("Otros", "otros_ganadores")],
            {"comprador": 45, "incumbente": 35, "otros_ganadores": 60}),
        "## Tu ranking actual frente al historial",
        md_table(res["ranking"], [
            ("#", "rank"), ("Código", "codigo"), ("Licitación", "nombre"), ("Base", "base"),
            ("Previas adj.", "previas_adjudicadas"), ("Ganadores previos", "ganadores_previos"),
            ("Oferentes prom.", "oferentes_promedio"), ("Monto mediano mismo tamaño (MM)", "monto_mediano_mm")],
            {"nombre": 50, "ganadores_previos": 80}),
    ]
    path.write_text("\n\n".join(parts) + "\n", encoding="utf-8")


def write_xlsx(path, res):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: skipping .xlsx output (pip install openpyxl)", file=sys.stderr)
        return False
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = [("Por tamaño", res["por_tamano"]), ("Competidores segmento", res["competidores_segmento"]),
              ("Proveedores", res["proveedores"]), ("Incumbentes", res["incumbentes"]),
              ("Ranking vs historial", res["ranking"]), ("Licitaciones históricas", res["licitaciones"])]
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        if not rows:
            continue
        cols = [k for k, v in rows[0].items() if not isinstance(v, (list, dict))]
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in rows:
            ws.append([r.get(k) for k in cols])
        for i, col in enumerate(cols, 1):
            wide = col in ("nombre", "comprador", "unidad", "compradores", "ejemplos", "ganadores_previos",
                           "otros_ganadores", "incumbente", "organismo", "ganadores")
            ws.column_dimensions[get_column_letter(i)].width = 55 if wide else 14
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return True


def write_html(path, res):
    # Self-contained page: the template plus the report embedded as JSON (tabs, search,
    # sorting and per-supplier bid history run in the browser, no server needed).
    template = pathlib.Path(__file__).with_name("competencia_template.html").read_text(encoding="utf-8")
    data = json.dumps(res, ensure_ascii=False).replace("</", "<\\/")
    path.write_text(template.replace("/*__DATA__*/", data), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--historico", default="data/historico.json")
    ap.add_argument("--profile", default="data/company_profile.json")
    ap.add_argument("--ranking", default="data/ranking_top100.json")
    ap.add_argument("--data", default="data/licitaciones.json")
    ap.add_argument("--out", default="data/competencia")
    ap.add_argument("--open", action="store_true", help="Open the .html in the browser")
    a = ap.parse_args()

    hist = json.loads(pathlib.Path(a.historico).read_text(encoding="utf-8"))
    profile = json.loads(pathlib.Path(a.profile).read_text(encoding="utf-8"))
    # historico.json also keeps loosely matching tenders for rank_history.py; competition uses the strict set.
    tenders = [t for t in hist["licitaciones"] if t.get("relevante", True)]
    sizes = {s.upper() for s in profile.get("tamanos", [])} or set(SIZE_ORDER)

    proveedores = supplier_stats(tenders, sizes)
    segmento = sorted((p for p in proveedores if p["ganadas_segmento"]),
                      key=lambda p: (-p["ganadas_segmento"], -p["monto_ganado_mm"]))

    ranking = []
    rank_path = pathlib.Path(a.ranking)
    if rank_path.exists():
        ranked = json.loads(rank_path.read_text(encoding="utf-8"))["ranking"]
        onu = {}
        data_path = pathlib.Path(a.data)
        if data_path.exists():
            for t in json.loads(data_path.read_text(encoding="utf-8"))["licitaciones"]:
                onu[t["codigo"]] = [i.get("codigo_onu", "") for i in t.get("items", []) if i.get("codigo_onu")]
        by_buyer = collections.defaultdict(list)
        for t in tenders:
            by_buyer[buyer_key(t["codigo"])].append(t)
        ranking = [benchmark(r, onu.get(r["codigo"], []), by_buyer, tenders) for r in ranked]

    flat = [{
        "codigo": t["codigo"], "nombre": t["nombre"], "organismo": t["organismo"], "region": t["region"],
        "tamano": t["tamano"], "estado": t["estado"], "fecha_publicacion": t["fecha_publicacion"],
        "fecha_adjudicacion": t["fecha_adjudicacion"], "n_oferentes": t["n_oferentes"],
        "monto_adjudicado_mm": mclp(awarded_amount(t)),
        "ganadores": ", ".join(o["nombre"] for o in winners(t)),
        "ofertas": [[o["rut"], o["nombre"], round(o["ofertado"]), round(o["adjudicado"]), o["gano"]]
                    for o in t["ofertas"]],
        "url": f"https://www.mercadopublico.cl/Procurement/Modules/RFB/DetailsAcquisition.aspx?idlicitacion={t['codigo']}",
    } for t in tenders]

    res = {
        "meta": {
            "empresa": profile.get("empresa", ""),
            "meses": f"{hist['meses'][0]} a {hist['meses'][-1]}",
            "generado_historico": hist["generado"],
            "licitaciones": len(tenders),
            "adjudicadas": sum(1 for t in tenders if winners(t)),
            "proveedores": len(proveedores),
            "tamanos_perfil": sorted(sizes, key=lambda s: SIZE_ORDER.index(s) if s in SIZE_ORDER else 9),
        },
        "por_tamano": summarize_sizes(tenders),
        "competidores_segmento": segmento,
        "proveedores": proveedores,
        "incumbentes": incumbents(tenders),
        "ranking": ranking,
        "licitaciones": flat,
    }
    out = pathlib.Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.with_suffix(".json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    write_md(out.with_suffix(".md"), res)
    write_html(out.with_suffix(".html"), res)
    files = [str(out.with_suffix(".json")), str(out.with_suffix(".md")), str(out.with_suffix(".html"))]
    if write_xlsx(out.with_suffix(".xlsx"), res):
        files.append(str(out.with_suffix(".xlsx")))

    with_history = sum(1 for r in ranking if r["previas_adjudicadas"])
    print(json.dumps({
        **res["meta"],
        "por_tamano": res["por_tamano"],
        "top10_segmento": [f"{p['ganadas_segmento']} ganadas · {p['nombre']} ({p['rut']})" for p in segmento[:10]],
        "incumbentes": len(res["incumbentes"]),
        "ranking_con_historial": f"{with_history}/{len(ranking)}",
        "salidas": files,
    }, ensure_ascii=False, indent=2))
    if a.open:
        sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2] / "mp-priorizar" / "scripts"))
        from rank import open_in_browser
        print(f"HTML abierto en {open_in_browser(out.with_suffix('.html'))}")


if __name__ == "__main__":
    main()
