"""Score past awarded tenders with the ranking's scoring system: jobs the company could have won, and who won them.

Each tender in data/historico.json that was awarded is scored by mp-priorizar's Ranker as it
stood on its publication day, so the time part reflects the days there were to prepare an offer
(publication to closing). The same filters apply as in the live ranking (exclusions, private
tenders, `solo_mis_regiones`, `solo_mis_tamanos`, `dias_minimos_preparacion`). Results are ordered
from the most recent award to the oldest, with the winners and every bid.

Writes data/historico_ranking.{json,csv,xlsx,html}.

Usage:
    python rank_history.py [--historico data/historico.json] [--profile data/company_profile.json]
                           [--min-puntaje 0] [--out data/historico_ranking] [--open]
"""
import argparse
import csv
import datetime as dt
import io
import json
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2] / "mp-priorizar" / "scripts"))
from rank import Ranker, job_summary, open_in_browser  # noqa: E402

COLUMNS = ["rank", "fecha_adjudicacion", "puntaje", "codigo", "nombre", "organismo", "region", "tamano",
           "monto_estimado", "monto_adjudicado", "ganadores", "n_oferentes", "ofertaste",
           "fecha_publicacion", "fecha_cierre", "dias_para_cierre", "motivos", "tipo", "url"]


def winners(t):
    return [o for o in t["ofertas"] if o["gano"]]


def award_date(t):
    # Fall back to publication when the award date is missing or a planned (future) date.
    d = t.get("fecha_adjudicacion") or ""
    return d if d and d <= dt.date.today().isoformat() else t["fecha_publicacion"]


def plain_rut(rut):
    return re.sub(r"[^0-9kK]", "", rut or "").upper()


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, rows):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: skipping .xlsx output (pip install openpyxl)", file=sys.stderr)
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Adjudicadas que calzan"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append([r.get(k) for k in COLUMNS])
        link = ws.cell(row=ws.max_row, column=COLUMNS.index("codigo") + 1)
        link.hyperlink = r["url"]
        link.font = Font(color="0563C1", underline="single")
    widths = {"nombre": 60, "organismo": 40, "ganadores": 45, "motivos": 70, "tipo": 35, "region": 28, "url": 30}
    for i, col in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 13)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=COLUMNS[c.column - 1] in ("nombre", "motivos", "ganadores"))
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return True


def write_html(path, payload):
    template = pathlib.Path(__file__).with_name("historico_template.html").read_text(encoding="utf-8")
    data = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    pathlib.Path(path).write_text(template.replace("/*__DATA__*/", data), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--historico", default="data/historico.json")
    ap.add_argument("--profile", default="data/company_profile.json")
    ap.add_argument("--min-puntaje", type=float, default=0)
    ap.add_argument("--out", default="data/historico_ranking")
    ap.add_argument("--open", action="store_true", help="open the HTML report in Chrome (or the default browser)")
    a = ap.parse_args()

    profile = json.loads(pathlib.Path(a.profile).read_text(encoding="utf-8"))
    hist = json.loads(pathlib.Path(a.historico).read_text(encoding="utf-8"))
    own_rut = plain_rut(profile.get("rut"))
    ranker = Ranker(profile, dt.datetime.now())

    rows, dropped, not_awarded = [], {}, {}
    for t in hist["licitaciones"]:
        won = winners(t)
        if not won:
            state = t["estado"].split(" (")[0] or "sin estado"
            not_awarded[state] = not_awarded.get(state, 0) + 1
            continue
        ranker.now = dt.datetime.fromisoformat(t["fecha_publicacion"])  # score as of publication day
        total, details = ranker.score(t)
        if total is None:
            dropped[details] = dropped.get(details, 0) + 1
            continue
        if total < a.min_puntaje:
            dropped["bajo el puntaje mínimo"] = dropped.get("bajo el puntaje mínimo", 0) + 1
            continue
        details["motivos"] = re.sub(r"cierra en (\d+) días", r"\1 días para ofertar", details["motivos"])
        rows.append({
            **{k: t.get(k, "") for k in ("codigo", "nombre", "organismo", "region", "tamano", "tipo", "url",
                                          "fecha_publicacion", "fecha_cierre", "estado")},
            "fecha_adjudicacion": award_date(t),
            "monto_estimado": round(t["monto_estimado"]) if t.get("monto_estimado", 0) > 1000 else None,
            "monto_adjudicado": round(sum(o["adjudicado"] for o in won)),
            "ganadores": ", ".join(o["nombre"] for o in won),
            "n_oferentes": t["n_oferentes"] or len(t["ofertas"]),
            "ofertaste": any(plain_rut(o["rut"]) == own_rut for o in t["ofertas"]) if own_rut else False,
            "ofertas": [[o["rut"], o["nombre"], round(o["ofertado"]), round(o["adjudicado"]), o["gano"]]
                        for o in t["ofertas"]],
            "n_items": len(t["items"]),
            **details,
            **job_summary(t),
        })
    rows.sort(key=lambda r: (r["fecha_adjudicacion"], r["puntaje"]), reverse=True)
    for i, r in enumerate(rows, 1):
        r["rank"] = i

    meta = {"titulo": f"Licitaciones adjudicadas {hist['meses'][0]} a {hist['meses'][-1]}",
            "descargado": hist.get("generado", ""), "generado": dt.datetime.now().isoformat(timespec="seconds")}
    awarded = sum(1 for t in hist["licitaciones"] if winners(t))
    stats = {"total": awarded, "ranked": len(rows), "dropped": dropped, "no_adjudicadas": not_awarded}
    payload = {"meta": meta, "stats": stats, "ranking": rows,
               "perfil": {k: profile.get(k) for k in ("empresa", "regiones", "tamanos")}}

    out = pathlib.Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.with_suffix(".json").write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    write_csv(out.with_suffix(".csv"), rows)
    write_html(out.with_suffix(".html"), payload)
    files = [str(out.with_suffix(e)) for e in (".json", ".csv", ".html")]
    if write_xlsx(out.with_suffix(".xlsx"), rows):
        files.append(str(out.with_suffix(".xlsx")))

    top = sorted(rows, key=lambda r: -r["puntaje"])[:10]
    print(json.dumps({
        "adjudicadas_en_historial": awarded,
        "calzan_con_perfil": len(rows),
        "descartadas": dropped,
        "salidas": files,
        "mejor_calce": [f"[{r['puntaje']}] {r['fecha_adjudicacion']} {r['codigo']} {r['nombre'][:55]} → {r['ganadores'][:40]}"
                        for r in top],
    }, ensure_ascii=False, indent=2))
    if a.open:
        opened = open_in_browser(out.with_suffix(".html"))
        print(f"HTML abierto en {opened}" if opened else "No se pudo abrir el HTML en un navegador")


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    main()
