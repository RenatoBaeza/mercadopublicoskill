"""Rank Mercado Público tenders against a company profile.

Score (0-100) =
    keywords  45  profile keywords found in title / items / categories / description
    category  25  UNSPSC (Código ONU) prefixes and category names matched by the items
    region    10  buyer region is one the company serves
    size      10  tender size bracket (L1/LE/LP/LQ/LR) is one the company can handle
    time      10  enough days left before the closing date

Tenders with no keyword or category match are not ranked. Tenders that already
closed, close before `dias_minimos_preparacion`, are private (unless
`incluir_privadas`), or match an exclusion are dropped.

The export has no amounts, so the estimated amount (Monto Total Estimado) of each ranked
tender is read from its page on mercadopublico.cl and cached next to the data file.

Usage:
    python rank.py --profile data/company_profile.json --data data/licitaciones.json \
                   [--top 100] [--out data/ranking_top100] [--now 2026-10-21T08:00] [--open] [--sin-montos]
"""
import argparse
import concurrent.futures
import csv
import datetime as dt
import io
import json
import math
import os
import pathlib
import re
import shutil
import subprocess
import sys
import unicodedata
import urllib.request
import webbrowser

W_KEYWORDS, W_CATEGORY, W_REGION, W_SIZE, W_TIME = 45, 25, 10, 10, 10
FIELD_WEIGHTS = {"nombre": 3.0, "items": 2.0, "categorias": 1.5, "descripcion": 1.0}
FIELD_LABELS = {"nombre": "título", "items": "ítems", "categorias": "categoría", "descripcion": "descripción"}
ONU_PREFIX_QUALITY = {8: 1.0, 6: 0.85, 4: 0.7, 2: 0.4}
# The export keeps some stale tenders with no closing date (some published in 2010);
# undated tenders older than this are treated as closed.
MAX_AGE_UNDATED_DAYS = 90
STOPWORDS = {"de", "del", "la", "las", "el", "los", "y", "e", "o", "u", "en", "para", "por",
             "con", "a", "al", "un", "una", "sus", "su", "the", "and", "of", "for"}
NATIONWIDE = {"", "todas", "todo chile", "nacional", "chile", "all"}


def norm(text):
    text = unicodedata.normalize("NFKD", str(text or "").lower())
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def stem(token):
    # Light prefix stem so "computador" matches "computadores", "limpieza" matches "limpiezas".
    return token if len(token) < 5 else token[:max(4, len(token) - 2)]


def phrase_stems(phrase):
    return [stem(t) for t in norm(phrase).split() if t not in STOPWORDS]


def matches(stems, tokens):
    # Short stems ("ti", "gps") must match a whole token, longer ones match as prefixes.
    return bool(stems) and all(
        (s in tokens) if len(s) < 4 else any(t.startswith(s) for t in tokens) for s in stems)


def tender_fields(t):
    items = t.get("items", [])
    return {
        "nombre": set(norm(t.get("nombre")).split()),
        "items": set(norm(" ".join(f"{i.get('item_descripcion', '')} {i.get('producto', '')}" for i in items)).split()),
        "categorias": set(norm(" ".join(f"{i.get('nivel1', '')} {i.get('nivel2', '')} {i.get('nivel3', '')}" for i in items)).split()),
        "descripcion": set(norm(t.get("descripcion")).split()),
    }


def parse_date(s):
    if not s:
        return None
    try:
        return dt.datetime.fromisoformat(s)
    except ValueError:
        return None


class Ranker:
    def __init__(self, profile, now):
        self.p = profile
        self.now = now
        self.kw = [(k, phrase_stems(k), 1.0) for k in profile.get("palabras_clave", [])]
        self.kw += [(k, phrase_stems(k), 0.5) for k in profile.get("palabras_clave_secundarias", [])]
        self.kw = [k for k in self.kw if k[1]]
        self.excl = [(k, phrase_stems(k)) for k in profile.get("excluir_palabras", [])]
        self.excl = [k for k in self.excl if k[1]]
        # Title-only exclusions: drop pure purchases ("adquisición de licencias") while keeping
        # service tenders that merely include such an item.
        self.excl_title = [(k, phrase_stems(k)) for k in profile.get("excluir_palabras_titulo", [])]
        self.excl_title = [k for k in self.excl_title if k[1]]
        self.excl_org = [norm(o) for o in profile.get("excluir_organismos", []) if norm(o)]
        self.onu = sorted({re.sub(r"\D", "", str(c)) for c in profile.get("codigos_onu", [])} - {""},
                          key=len, reverse=True)
        self.cats = [(c, phrase_stems(c)) for c in profile.get("categorias", [])]
        self.cats = [c for c in self.cats if c[1]]
        regions = [norm(r) for r in profile.get("regiones", [])]
        self.nationwide = not regions or any(r in NATIONWIDE for r in regions)
        self.regions = [r.replace("region", "").strip() for r in regions]
        self.sizes = {s.upper() for s in profile.get("tamanos", [])}
        self.min_days = float(profile.get("dias_minimos_preparacion", 3))
        self.include_private = bool(profile.get("incluir_privadas", False))
        self.only_my_regions = bool(profile.get("solo_mis_regiones", False))
        self.only_my_sizes = bool(profile.get("solo_mis_tamanos", False))

    def item_category_quality(self, item):
        best, why = 0.0, ""
        code = item.get("codigo_onu", "")
        for prefix in self.onu:
            if code.startswith(prefix):
                q = ONU_PREFIX_QUALITY.get(len(prefix), 0.5)
                if q > best:
                    best, why = q, f"ONU {prefix}"
                break  # prefixes sorted longest first
        if self.cats:
            levels = (("nivel3", 0.8), ("producto", 0.8), ("nivel2", 0.6), ("nivel1", 0.4))
            for name, stems in self.cats:
                for field, q in levels:
                    if q > best and matches(stems, set(norm(item.get(field)).split())):
                        best, why = q, f"cat. {name}"
        return best, why

    def score(self, t):
        """Returns (score, details) or (None, reason) when the tender is filtered out."""
        cierre = parse_date(t.get("fecha_cierre"))
        days_left = (cierre - self.now).total_seconds() / 86400 if cierre else None
        if days_left is not None and days_left < 0:
            return None, "cerrada"
        if cierre is None:
            publicada = parse_date(t.get("fecha_publicacion"))
            if publicada is None or (self.now - publicada).days > MAX_AGE_UNDATED_DAYS:
                return None, "sin fecha de cierre y antigua"
        if days_left is not None and days_left < self.min_days:
            return None, "cierra muy pronto"
        if t.get("privada") and not self.include_private:
            return None, "privada"
        if self.excl_org and any(o in norm(t.get("organismo")) for o in self.excl_org):
            return None, "organismo excluido"

        fields = tender_fields(t)
        for word, stems in self.excl:
            if any(matches(stems, fields[f]) for f in ("nombre", "items")):
                return None, f"excluida por '{word}'"
        for word, stems in self.excl_title:
            if matches(stems, fields["nombre"]):
                return None, f"excluida por '{word}' (título)"

        region_ok = self.nationwide or any(r and r in norm(t.get("region")) for r in self.regions)
        if self.only_my_regions and not region_ok:
            return None, "fuera de región"

        size = t.get("tamano", "")
        size_known = size not in ("LS", "OBRA", "")
        if self.only_my_sizes and self.sizes and size_known and size not in self.sizes:
            return None, "fuera de tamaño"

        # Keywords: each keyword counts once, at its best-weighted field.
        raw, hits = 0.0, []
        for word, stems, weight in self.kw:
            best_field = max((f for f in FIELD_WEIGHTS if matches(stems, fields[f])),
                             key=FIELD_WEIGHTS.get, default=None)
            if best_field:
                raw += weight * FIELD_WEIGHTS[best_field]
                hits.append(f"{word} ({FIELD_LABELS[best_field]})")
        kw_score = W_KEYWORDS * (1 - math.exp(-raw / 4))

        # Category: best item match, scaled by the share of items that match at all.
        items = t.get("items") or [{}]
        qualities = [self.item_category_quality(i) for i in items]
        best_q, best_why = max(qualities, key=lambda x: x[0])
        share = sum(1 for q, _ in qualities if q > 0) / len(items)
        cat_score = W_CATEGORY * best_q * (0.6 + 0.4 * share) if best_q else 0.0

        if kw_score == 0 and cat_score == 0:
            return None, "sin coincidencias"

        if not self.sizes or size in self.sizes:
            size_score = W_SIZE
        elif not size_known:
            size_score = W_SIZE / 2  # size unknown from the tender type
        else:
            size_score = 0.0

        if days_left is None:
            time_score = W_TIME / 2
        else:
            time_score = W_TIME * min(1.0, (days_left - self.min_days + 1) / 7)

        region_score = W_REGION if region_ok else 0.0
        total = kw_score + cat_score + region_score + size_score + time_score

        reasons = []
        if hits:
            reasons.append("palabras: " + ", ".join(hits))
        if best_q:
            reasons.append(f"{best_why} ({share:.0%} de los ítems)")
        reasons.append("región ✓" if region_ok else "otra región")
        reasons.append(f"tamaño {size or '?'} {'✓' if size_score == W_SIZE else '~' if size_score else '✗'}")
        if days_left is not None:
            reasons.append(f"cierra en {days_left:.0f} días")
        else:
            reasons.append("sin fecha de cierre: verificar en el portal")
        return round(total, 1), {
            "puntaje": round(total, 1),
            "p_palabras": round(kw_score, 1),
            "p_categoria": round(cat_score, 1),
            "p_region": round(region_score, 1),
            "p_tamano": round(size_score, 1),
            "p_plazo": round(time_score, 1),
            "dias_para_cierre": round(days_left, 1) if days_left is not None else None,
            "motivos": " | ".join(reasons),
        }


BOILERPLATE = re.compile(r"^(de acuerdo|segun|según|conforme|ver|indicad|establecid|detallad|las? (bases|especificaciones))", re.I)
MAX_DELIVERABLES = 30


def readable(text):
    # Many tenders are typed in ALL CAPS; sentence-case them so the summary reads normally.
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    letters = [c for c in text if c.isalpha()]
    if letters and sum(c.isupper() for c in letters) / len(letters) > 0.6:
        text = text.lower()
        text = re.sub(r"(^|[.!?]\s+)([a-záéíóúñ])", lambda m: m.group(1) + m.group(2).upper(), text)
    return text


def job_summary(t):
    """What the buyer wants: the description and the list of deliverables (what, how many)."""
    deliverables, seen = [], set()
    for i in t.get("items", []):
        what = readable(i.get("item_descripcion"))
        if not what or BOILERPLATE.match(what) or norm(what) == norm(t.get("descripcion")):
            what = readable(i.get("producto")) or what
        key = norm(what)
        if not key or key in seen:
            continue
        seen.add(key)
        qty, unit = i.get("cantidad"), i.get("unidad") or ""
        amount = ""
        if isinstance(qty, (int, float)) and (qty != 1 or unit.lower() not in ("unidad", "")):
            amount = f"{qty:g} {unit}".strip()
        deliverables.append({"que": what, "cantidad": amount})
    return {"descripcion": readable(t.get("descripcion")),
            "entregables": deliverables[:MAX_DELIVERABLES],
            "entregables_extra": max(0, len(deliverables) - MAX_DELIVERABLES)}


# Amounts published on each tender's page (the export doesn't carry them).
AMOUNT_SPANS = {"moneda": "lblFicha1Moneda", "monto_estimado": "lblFicha7MontoEstimado",
                "monto_base": "lblFicha7Estimacion"}
CURRENCY_CODES = {"peso chileno": "CLP", "unidad de fomento": "UF", "dolar": "USD", "euro": "EUR",
                  "unidad tributaria mensual": "UTM"}


def parse_amount(s):
    s = re.sub(r"[^\d,.]", "", s or "")
    if "," in s:  # "15.000.000,50" -> 15000000.50
        s = s.replace(".", "").replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return None
    return int(n) if n.is_integer() else n


def fetch_amount(url):
    """Reads currency, estimated amount and its basis from a tender page. Amount is None when not public."""
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (mp-priorizar skill)"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        page = resp.read().decode("utf-8", "replace")
    found = {}
    for key, span in AMOUNT_SPANS.items():
        m = re.search(r'id="' + span + r'"[^>]*>(.*?)</span>', page, re.S)
        found[key] = re.sub(r"<[^>]+>|\s+", " ", m.group(1)).strip() if m else ""
    currency = norm(found["moneda"])
    found["moneda"] = next((c for name, c in CURRENCY_CODES.items() if name in currency), found["moneda"])
    found["monto_estimado"] = parse_amount(found["monto_estimado"])
    return found


def add_amounts(rows, cache_path, workers=8):
    """Adds monto_estimado / moneda / monto_base to each row, fetching only tenders not in the cache."""
    try:
        with open(cache_path, encoding="utf-8") as f:
            cache = json.load(f)
    except (OSError, ValueError):
        cache = {}
    todo = [r for r in rows if r["codigo"] not in cache]
    failed = 0
    with concurrent.futures.ThreadPoolExecutor(workers) as ex:
        futures = {ex.submit(fetch_amount, r["url"]): r["codigo"] for r in todo}
        for fut in concurrent.futures.as_completed(futures):
            try:
                cache[futures[fut]] = fut.result()
            except Exception:  # noqa: BLE001 - a missing amount is not fatal; retried next run
                failed += 1
    if todo:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=1)
    for r in rows:
        r.update(cache.get(r["codigo"], {}))
    return {"consultadas": len(todo) - failed, "en_cache": len(rows) - len(todo), "fallidas": failed,
            "con_monto": sum(1 for r in rows if r.get("monto_estimado") is not None)}


OUTPUT_COLUMNS = ["rank", "puntaje", "codigo", "nombre", "organismo", "region", "tamano",
                  "monto_estimado", "moneda", "tipo",
                  "fecha_cierre", "dias_para_cierre", "n_items", "motivos", "p_palabras",
                  "p_categoria", "p_region", "p_tamano", "p_plazo", "url"]


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, rows, profile, meta):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: skipping .xlsx output (pip install openpyxl)", file=sys.stderr)
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Top licitaciones"
    ws.append(OUTPUT_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append([r.get(k) for k in OUTPUT_COLUMNS])
        link = ws.cell(row=ws.max_row, column=OUTPUT_COLUMNS.index("codigo") + 1)
        link.hyperlink = r["url"]
        link.font = Font(color="0563C1", underline="single")
    widths = {"nombre": 60, "organismo": 40, "motivos": 70, "tipo": 35, "region": 28, "url": 30}
    for i, col in enumerate(OUTPUT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 13)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=OUTPUT_COLUMNS[c.column - 1] in ("nombre", "motivos"))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Perfil y fuente")
    info.append(["Fuente", meta.get("titulo", "")])
    info.append(["Descargado", meta.get("descargado", "")])
    info.append(["Ranking generado", meta.get("generado", "")])
    for k, v in profile.items():
        info.append([k, ", ".join(map(str, v)) if isinstance(v, list) else str(v)])
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 100
    wb.save(path)
    return True


def write_md(path, rows, profile, meta, stats):
    def cell(s, n=None):
        s = str(s or "").replace("|", "/").replace("\n", " ")
        return s if n is None or len(s) <= n else s[: n - 1] + "…"

    lines = [
        f"# Top {len(rows)} licitaciones para {profile.get('empresa', 'la empresa')}",
        "",
        f"- Fuente: {meta.get('titulo', '')} (descargado {meta.get('descargado', '')})",
        f"- Licitaciones evaluadas: {stats['total']} · con coincidencias: {stats['ranked']}",
        f"- Descartadas: " + ", ".join(f"{k}: {v}" for k, v in sorted(stats["dropped"].items(), key=lambda x: -x[1])),
        "",
        "| # | Puntaje | Código | Nombre | Organismo | Región | Tamaño | Cierre | Motivos |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for r in rows:
        lines.append(
            f"| {r['rank']} | {r['puntaje']} | [{r['codigo']}]({r['url']}) | {cell(r['nombre'], 70)} | "
            f"{cell(r['organismo'], 40)} | {cell(r['region'].replace('Región ', ''), 25)} | {r['tamano']} | "
            f"{cell((r['fecha_cierre'] or '')[:16].replace('T', ' '))} | {cell(r['motivos'])} |"
        )
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def write_html(path, rows, profile, meta, stats):
    # Self-contained page: the template plus the ranking embedded as JSON (filters, sorting,
    # job summaries and a starred shortlist run in the browser, no server needed).
    template = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ranking_template.html")
    with open(template, encoding="utf-8") as f:
        html = f.read()
    payload = {"meta": meta, "stats": stats, "ranking": rows,
               "perfil": {k: profile.get(k) for k in ("empresa", "regiones", "tamanos")}}
    data = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html.replace("/*__DATA__*/", data))


def open_in_browser(path):
    """Opens the file in Chrome if installed, else in the default browser. Returns which one."""
    url = pathlib.Path(path).resolve().as_uri()
    if sys.platform == "win32":
        candidates = [os.path.join(os.environ[v], "Google", "Chrome", "Application", "chrome.exe")
                      for v in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA") if os.environ.get(v)]
    elif sys.platform == "darwin":
        candidates = ["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"]
    else:
        candidates = [shutil.which(n) for n in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser")]
    for exe in candidates:
        if exe and os.path.isfile(exe):
            subprocess.Popen([exe, url], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            return "chrome"
    return "navegador predeterminado" if webbrowser.open(url) else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--profile", default="data/company_profile.json")
    ap.add_argument("--data", default="data/licitaciones.json")
    ap.add_argument("--top", type=int, default=100)
    ap.add_argument("--out", default="data/ranking_top100", help="output path prefix (no extension)")
    ap.add_argument("--now", help="reference time, ISO format (default: current local time)")
    ap.add_argument("--open", action="store_true", help="open the HTML report in Chrome (or the default browser)")
    ap.add_argument("--sin-montos", action="store_true",
                    help="don't fetch estimated amounts from mercadopublico.cl (offline / quick re-runs)")
    args = ap.parse_args()

    with open(args.profile, encoding="utf-8") as f:
        profile = json.load(f)
    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)
    now = dt.datetime.fromisoformat(args.now) if args.now else dt.datetime.now()

    ranker = Ranker(profile, now)
    if not ranker.kw and not ranker.onu and not ranker.cats:
        sys.exit("ERROR: profile has no palabras_clave, codigos_onu or categorias to match on")

    scored, dropped = [], {}
    tenders = data.get("licitaciones", [])
    for t in tenders:
        total, details = ranker.score(t)
        if total is None:
            dropped[details] = dropped.get(details, 0) + 1
            continue
        scored.append({
            **{k: t.get(k, "") for k in ("codigo", "nombre", "organismo", "region", "tamano", "tipo",
                                          "fecha_cierre", "url")},
            "n_items": len(t.get("items", [])),
            **details,
        })
    scored.sort(key=lambda r: (-r["puntaje"], r["dias_para_cierre"] or 0))
    top = scored[: args.top]
    by_code = {t.get("codigo"): t for t in tenders}
    for i, r in enumerate(top, 1):
        r["rank"] = i
        r.update(job_summary(by_code[r["codigo"]]))
    amounts = None
    if not args.sin_montos:
        cache = os.path.join(os.path.dirname(os.path.abspath(args.data)), "montos_cache.json")
        amounts = add_amounts(top, cache)

    meta = {"titulo": data.get("titulo", ""), "descargado": data.get("descargado", ""),
            "generado": now.replace(microsecond=0).isoformat()}
    stats = {"total": len(tenders), "ranked": len(scored), "dropped": dropped}
    write_csv(args.out + ".csv", top)
    with open(args.out + ".json", "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "stats": stats, "ranking": top}, f, ensure_ascii=False, indent=1)
    write_md(args.out + ".md", top, profile, meta, stats)
    write_html(args.out + ".html", top, profile, meta, stats)
    has_xlsx = write_xlsx(args.out + ".xlsx", top, profile, meta)

    print(json.dumps({
        "evaluadas": len(tenders),
        "con_coincidencias": len(scored),
        "descartadas": dropped,
        "top": len(top),
        **({"montos": amounts} if amounts else {}),
        "salidas": [args.out + ext for ext in ([".xlsx"] if has_xlsx else []) + [".csv", ".json", ".md", ".html"]],
        "top10": [f"{r['rank']}. [{r['puntaje']}] {r['codigo']} {r['nombre'][:70]}" for r in top[:10]],
    }, ensure_ascii=False, indent=2))
    if args.open:
        opened = open_in_browser(args.out + ".html")
        print(f"HTML abierto en {opened}" if opened else "No se pudo abrir el HTML en un navegador")


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    main()
