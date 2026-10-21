"""Rank Mercado Público tenders against a company profile.

Score (0-100) =
    keywords  45  profile keywords found in title / items / categories / description
    category  25  UNSPSC (Código ONU) prefixes and category names matched by the items
    region    10  buyer region is one the company serves
    size      10  tender size bracket (L1/LE/LP/LQ/LR) is one the company can handle
    time      10  enough days left before the closing date

Tenders with no keyword or category match are not ranked. Tenders that already
closed, close before `dias_minimos_preparacion`, are private (unless
`incluir_privadas`), or match an exclusion are dropped.

Usage:
    python rank.py --profile data/company_profile.json --data data/licitaciones.json \
                   [--top 100] [--out data/ranking_top100] [--now 2026-10-21T08:00]
"""
import argparse
import csv
import datetime as dt
import io
import json
import math
import re
import sys
import unicodedata

W_KEYWORDS, W_CATEGORY, W_REGION, W_SIZE, W_TIME = 45, 25, 10, 10, 10
FIELD_WEIGHTS = {"nombre": 3.0, "items": 2.0, "categorias": 1.5, "descripcion": 1.0}
FIELD_LABELS = {"nombre": "título", "items": "ítems", "categorias": "categoría", "descripcion": "descripción"}
ONU_PREFIX_QUALITY = {8: 1.0, 6: 0.85, 4: 0.7, 2: 0.4}
STOPWORDS = {"de", "del", "la", "las", "el", "los", "y", "e", "o", "u", "en", "para", "por",
             "con", "a", "al", "un", "una", "sus", "su", "the", "and", "of", "for"}
NATIONWIDE = {"", "todas", "todo chile", "nacional", "chile", "all"}


def norm(text):
    text = unicodedata.normalize("NFKD", str(text or "").lower())
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def stem(token):
    # Light prefix stem so "computador" matches "computadores", "limpieza" matches "limpiezas".
    return token if len(token) < 5 else token[:max(4, len(token) - 2)]


def phrase_stems(phrase):
    return [stem(t) for t in norm(phrase).split() if t not in STOPWORDS]


def matches(stems, tokens):
    # Short stems ("ti", "gps") must match a whole token, longer ones match as prefixes.
    return bool(stems) and all(
        (s in tokens) if len(s) < 4 else any(t.startswith(s) for t in tokens) for s in stems)


def tender_fields(t):
    items = t.get("items", [])
    return {
        "nombre": set(norm(t.get("nombre")).split()),
        "items": set(norm(" ".join(f"{i.get('item_descripcion', '')} {i.get('producto', '')}" for i in items)).split()),
        "categorias": set(norm(" ".join(f"{i.get('nivel1', '')} {i.get('nivel2', '')} {i.get('nivel3', '')}" for i in items)).split()),
        "descripcion": set(norm(t.get("descripcion")).split()),
    }


def parse_date(s):
    if not s:
        return None
    try:
        return dt.datetime.fromisoformat(s)
    except ValueError:
        return None


class Ranker:
    def __init__(self, profile, now):
        self.p = profile
        self.now = now
        self.kw = [(k, phrase_stems(k), 1.0) for k in profile.get("palabras_clave", [])]
        self.kw += [(k, phrase_stems(k), 0.5) for k in profile.get("palabras_clave_secundarias", [])]
        self.kw = [k for k in self.kw if k[1]]
        self.excl = [(k, phrase_stems(k)) for k in profile.get("excluir_palabras", [])]
        self.excl = [k for k in self.excl if k[1]]
        # Title-only exclusions: drop pure purchases ("adquisición de licencias") while keeping
        # service tenders that merely include such an item.
        self.excl_title = [(k, phrase_stems(k)) for k in profile.get("excluir_palabras_titulo", [])]
        self.excl_title = [k for k in self.excl_title if k[1]]
        self.excl_org = [norm(o) for o in profile.get("excluir_organismos", []) if norm(o)]
        self.onu = sorted({re.sub(r"\D", "", str(c)) for c in profile.get("codigos_onu", [])} - {""},
                          key=len, reverse=True)
        self.cats = [(c, phrase_stems(c)) for c in profile.get("categorias", [])]
        self.cats = [c for c in self.cats if c[1]]
        regions = [norm(r) for r in profile.get("regiones", [])]
        self.nationwide = not regions or any(r in NATIONWIDE for r in regions)
        self.regions = [r.replace("region", "").strip() for r in regions]
        self.sizes = {s.upper() for s in profile.get("tamanos", [])}
        self.min_days = float(profile.get("dias_minimos_preparacion", 3))
        self.include_private = bool(profile.get("incluir_privadas", False))
        self.only_my_regions = bool(profile.get("solo_mis_regiones", False))

    def item_category_quality(self, item):
        best, why = 0.0, ""
        code = item.get("codigo_onu", "")
        for prefix in self.onu:
            if code.startswith(prefix):
                q = ONU_PREFIX_QUALITY.get(len(prefix), 0.5)
                if q > best:
                    best, why = q, f"ONU {prefix}"
                break  # prefixes sorted longest first
        if self.cats:
            levels = (("nivel3", 0.8), ("producto", 0.8), ("nivel2", 0.6), ("nivel1", 0.4))
            for name, stems in self.cats:
                for field, q in levels:
                    if q > best and matches(stems, set(norm(item.get(field)).split())):
                        best, why = q, f"cat. {name}"
        return best, why

    def score(self, t):
        """Returns (score, details) or (None, reason) when the tender is filtered out."""
        cierre = parse_date(t.get("fecha_cierre"))
        days_left = (cierre - self.now).total_seconds() / 86400 if cierre else None
        if days_left is not None and days_left < 0:
            return None, "cerrada"
        if days_left is not None and days_left < self.min_days:
            return None, "cierra muy pronto"
        if t.get("privada") and not self.include_private:
            return None, "privada"
        if self.excl_org and any(o in norm(t.get("organismo")) for o in self.excl_org):
            return None, "organismo excluido"

        fields = tender_fields(t)
        for word, stems in self.excl:
            if any(matches(stems, fields[f]) for f in ("nombre", "items")):
                return None, f"excluida por '{word}'"
        for word, stems in self.excl_title:
            if matches(stems, fields["nombre"]):
                return None, f"excluida por '{word}' (título)"

        region_ok = self.nationwide or any(r and r in norm(t.get("region")) for r in self.regions)
        if self.only_my_regions and not region_ok:
            return None, "fuera de región"

        # Keywords: each keyword counts once, at its best-weighted field.
        raw, hits = 0.0, []
        for word, stems, weight in self.kw:
            best_field = max((f for f in FIELD_WEIGHTS if matches(stems, fields[f])),
                             key=FIELD_WEIGHTS.get, default=None)
            if best_field:
                raw += weight * FIELD_WEIGHTS[best_field]
                hits.append(f"{word} ({FIELD_LABELS[best_field]})")
        kw_score = W_KEYWORDS * (1 - math.exp(-raw / 4))

        # Category: best item match, scaled by the share of items that match at all.
        items = t.get("items") or [{}]
        qualities = [self.item_category_quality(i) for i in items]
        best_q, best_why = max(qualities, key=lambda x: x[0])
        share = sum(1 for q, _ in qualities if q > 0) / len(items)
        cat_score = W_CATEGORY * best_q * (0.6 + 0.4 * share) if best_q else 0.0

        if kw_score == 0 and cat_score == 0:
            return None, "sin coincidencias"

        size = t.get("tamano", "")
        if not self.sizes or size in self.sizes:
            size_score = W_SIZE
        elif size in ("LS", "OBRA", ""):
            size_score = W_SIZE / 2  # size unknown from the tender type
        else:
            size_score = 0.0

        if days_left is None:
            time_score = W_TIME / 2
        else:
            time_score = W_TIME * min(1.0, (days_left - self.min_days + 1) / 7)

        region_score = W_REGION if region_ok else 0.0
        total = kw_score + cat_score + region_score + size_score + time_score

        reasons = []
        if hits:
            reasons.append("palabras: " + ", ".join(hits))
        if best_q:
            reasons.append(f"{best_why} ({share:.0%} de los ítems)")
        reasons.append("región ✓" if region_ok else "otra región")
        reasons.append(f"tamaño {size or '?'} {'✓' if size_score == W_SIZE else '~' if size_score else '✗'}")
        if days_left is not None:
            reasons.append(f"cierra en {days_left:.0f} días")
        return round(total, 1), {
            "puntaje": round(total, 1),
            "p_palabras": round(kw_score, 1),
            "p_categoria": round(cat_score, 1),
            "p_region": round(region_score, 1),
            "p_tamano": round(size_score, 1),
            "p_plazo": round(time_score, 1),
            "dias_para_cierre": round(days_left, 1) if days_left is not None else None,
            "motivos": " | ".join(reasons),
        }


OUTPUT_COLUMNS = ["rank", "puntaje", "codigo", "nombre", "organismo", "region", "tamano", "tipo",
                  "fecha_cierre", "dias_para_cierre", "n_items", "motivos", "p_palabras",
                  "p_categoria", "p_region", "p_tamano", "p_plazo", "url"]


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, rows, profile, meta):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: skipping .xlsx output (pip install openpyxl)", file=sys.stderr)
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Top licitaciones"
    ws.append(OUTPUT_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append([r.get(k) for k in OUTPUT_COLUMNS])
        link = ws.cell(row=ws.max_row, column=OUTPUT_COLUMNS.index("codigo") + 1)
        link.hyperlink = r["url"]
        link.font = Font(color="0563C1", underline="single")
    widths = {"nombre": 60, "organismo": 40, "motivos": 70, "tipo": 35, "region": 28, "url": 30}
    for i, col in enumerate(OUTPUT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 13)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=OUTPUT_COLUMNS[c.column - 1] in ("nombre", "motivos"))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Perfil y fuente")
    info.append(["Fuente", meta.get("titulo", "")])
    info.append(["Descargado", meta.get("descargado", "")])
    info.append(["Ranking generado", meta.get("generado", "")])
    for k, v in profile.items():
        info.append([k, ", ".join(map(str, v)) if isinstance(v, list) else str(v)])
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 100
    wb.save(path)
    return True


def write_md(path, rows, profile, meta, stats):
    def cell(s, n=None):
        s = str(s or "").replace("|", "/").replace("\n", " ")
        return s if n is None or len(s) <= n else s[: n - 1] + "…"

    lines = [
        f"# Top {len(rows)} licitaciones para {profile.get('empresa', 'la empresa')}",
        "",
        f"- Fuente: {meta.get('titulo', '')} (descargado {meta.get('descargado', '')})",
        f"- Licitaciones evaluadas: {stats['total']} · con coincidencias: {stats['ranked']}",
        f"- Descartadas: " + ", ".join(f"{k}: {v}" for k, v in sorted(stats["dropped"].items(), key=lambda x: -x[1])),
        "",
        "| # | Puntaje | Código | Nombre | Organismo | Región | Tamaño | Cierre | Motivos |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for r in rows:
        lines.append(
            f"| {r['rank']} | {r['puntaje']} | [{r['codigo']}]({r['url']}) | {cell(r['nombre'], 70)} | "
            f"{cell(r['organismo'], 40)} | {cell(r['region'].replace('Región ', ''), 25)} | {r['tamano']} | "
            f"{cell((r['fecha_cierre'] or '')[:16].replace('T', ' '))} | {cell(r['motivos'])} |"
        )
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--profile", default="data/company_profile.json")
    ap.add_argument("--data", default="data/licitaciones.json")
    ap.add_argument("--top", type=int, default=100)
    ap.add_argument("--out", default="data/ranking_top100", help="output path prefix (no extension)")
    ap.add_argument("--now", help="reference time, ISO format (default: current local time)")
    args = ap.parse_args()

    with open(args.profile, encoding="utf-8") as f:
        profile = json.load(f)
    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)
    now = dt.datetime.fromisoformat(args.now) if args.now else dt.datetime.now()

    ranker = Ranker(profile, now)
    if not ranker.kw and not ranker.onu and not ranker.cats:
        sys.exit("ERROR: profile has no palabras_clave, codigos_onu or categorias to match on")

    scored, dropped = [], {}
    tenders = data.get("licitaciones", [])
    for t in tenders:
        total, details = ranker.score(t)
        if total is None:
            dropped[details] = dropped.get(details, 0) + 1
            continue
        scored.append({
            **{k: t.get(k, "") for k in ("codigo", "nombre", "organismo", "region", "tamano", "tipo",
                                          "fecha_cierre", "url")},
            "n_items": len(t.get("items", [])),
            **details,
        })
    scored.sort(key=lambda r: (-r["puntaje"], r["dias_para_cierre"] or 0))
    top = scored[: args.top]
    for i, r in enumerate(top, 1):
        r["rank"] = i

    meta = {"titulo": data.get("titulo", ""), "descargado": data.get("descargado", ""),
            "generado": now.replace(microsecond=0).isoformat()}
    stats = {"total": len(tenders), "ranked": len(scored), "dropped": dropped}
    write_csv(args.out + ".csv", top)
    with open(args.out + ".json", "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "stats": stats, "ranking": top}, f, ensure_ascii=False, indent=1)
    write_md(args.out + ".md", top, profile, meta, stats)
    has_xlsx = write_xlsx(args.out + ".xlsx", top, profile, meta)

    print(json.dumps({
        "evaluadas": len(tenders),
        "con_coincidencias": len(scored),
        "descartadas": dropped,
        "top": len(top),
        "salidas": [args.out + ext for ext in ([".xlsx"] if has_xlsx else []) + [".csv", ".json", ".md"]],
        "top10": [f"{r['rank']}. [{r['puntaje']}] {r['codigo']} {r['nombre'][:70]}" for r in top[:10]],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    main()
